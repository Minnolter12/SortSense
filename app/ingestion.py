"""
FileMind document ingestion.

Converts a file on disk into a NormalizedDocument.
Supports:
  - PDF (via PyMuPDF with pypdf/PyPDF2 fallback)
  - Images (.png, .jpg, .jpeg, .webp, .bmp, .tiff via pytesseract OCR)
  - Office Documents (.docx, .xlsx, .pptx via python-docx/openpyxl/python-pptx with stdlib XML fallback)
  - Plain text, Markdown, Data & Code files (.txt, .md, .csv, .json, .py, .js, .ts, .html, .xml, .yaml, .yml, .log, .ini, .toml, .sql, .sh)

All errors are represented in the returned NormalizedDocument rather than
propagated as exceptions, except for programming errors (wrong argument types).
No dependency on the watcher, AI, Ollama, or database layers.
"""

from __future__ import annotations

import logging
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from app.document import NormalizedDocument

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Content-type mapping
# ---------------------------------------------------------------------------

_CONTENT_TYPES: dict[str, str] = {
    # Documents
    "pdf":  "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    # Text / Markdown / Code / Data
    "txt":  "text/plain",
    "md":   "text/markdown",
    "csv":  "text/csv",
    "json": "application/json",
    "py":   "text/x-python",
    "js":   "text/javascript",
    "ts":   "text/typescript",
    "html": "text/html",
    "xml":  "application/xml",
    "yaml": "application/x-yaml",
    "yml":  "application/x-yaml",
    "log":  "text/plain",
    "ini":  "text/plain",
    "toml": "application/toml",
    "sql":  "application/sql",
    "sh":   "text/x-shellscript",
    # Images
    "png":  "image/png",
    "jpg":  "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
    "bmp":  "image/bmp",
    "tiff": "image/tiff",
}

_TEXT_EXTENSIONS = {
    "txt", "md", "csv", "json", "py", "js", "ts",
    "html", "xml", "yaml", "yml", "log", "ini", "toml", "sql", "sh",
}

_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "bmp", "tiff"}


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _base_fields(path: Path) -> dict:
    """Return the fields that are always the same regardless of outcome."""
    ext = path.suffix.lstrip(".").lower()
    return {
        "file_path":    str(path),
        "file_name":    path.name,
        "extension":    ext,
        "content_type": _CONTENT_TYPES.get(ext, "application/octet-stream"),
    }


def _make_doc(path: Path, text: str, page_count: int | None, status: str) -> NormalizedDocument:
    return NormalizedDocument(
        **_base_fields(path),
        text=text,
        character_count=len(text),
        page_count=page_count,
        extraction_status=status,
    )


def _ingest_pdf(path: Path) -> NormalizedDocument:
    """Extract text from a PDF using PyMuPDF (primary) or pypdf/PyPDF2 (fallback)."""
    try:
        import pymupdf
        try:
            doc = pymupdf.open(str(path))
        except Exception as exc:
            logger.warning("Could not open PDF %s: %s", path, exc)
            return _make_doc(path, text="", page_count=None, status="error")

        try:
            pages: list[str] = []
            for page in doc:
                pages.append(page.get_text())
            text = "\n".join(pages)
            page_count = len(doc)
        except Exception as exc:
            logger.warning("Error reading PDF pages in %s: %s", path, exc)
            return _make_doc(path, text="", page_count=None, status="error")
        finally:
            doc.close()

        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=page_count, status=status)
    except ImportError:
        pass

    # Fallback to pypdf / PyPDF2 if PyMuPDF is not installed
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader  # type: ignore[no-redef]

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages)
        page_count = len(reader.pages)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=page_count, status=status)
    except Exception as exc:
        logger.warning("Could not read PDF %s via fallback: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")


def _ingest_text(path: Path) -> NormalizedDocument:
    """Read a plain-text, Markdown, CSV, JSON, or source code file as UTF-8."""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Fall back to latin-1 for files that are not valid UTF-8
        try:
            text = path.read_text(encoding="latin-1")
        except Exception as exc:
            logger.warning("Could not decode %s: %s", path, exc)
            return _make_doc(path, text="", page_count=None, status="error")
    except Exception as exc:
        logger.warning("Could not read %s: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")

    status = "ok" if text.strip() else "empty"
    return _make_doc(path, text=text, page_count=None, status=status)


def _ingest_image(path: Path) -> NormalizedDocument:
    """Extract text from an image file using pytesseract OCR."""
    try:
        from PIL import Image
        img = Image.open(path)
        img.verify()
        # Re-open after verify() so pytesseract can read pixel data
        img = Image.open(path)
    except ImportError:
        # If Pillow is unavailable, check basic file readability
        try:
            raw = path.read_bytes()
            if not raw:
                return _make_doc(path, text="", page_count=None, status="empty")
            img = None
        except Exception as exc:
            logger.warning("Could not read image %s: %s", path, exc)
            return _make_doc(path, text="", page_count=None, status="error")
    except Exception as exc:
        logger.warning("Invalid or corrupt image %s: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")

    try:
        import pytesseract
        target = img if img is not None else str(path)
        text = pytesseract.image_to_string(target)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=1, status=status)
    except Exception as exc:
        # Tesseract binary not installed or OCR unavailable — return "empty" so
        # the classifier can still categorize based on filename/image metadata
        logger.info("OCR unavailable or produced no text for %s: %s", path, exc)
        return _make_doc(path, text="", page_count=1, status="empty")


def _ingest_docx(path: Path) -> NormalizedDocument:
    """Extract text from a Word (.docx) document."""
    try:
        import docx  # python-docx
        doc = docx.Document(str(path))
        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        text = "\n".join(parts)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=None, status=status)
    except ImportError:
        pass
    except Exception as exc:
        logger.debug("python-docx failed on %s (%s), trying XML fallback", path, exc)

    # Stdlib zipfile + XML fallback for .docx
    try:
        with zipfile.ZipFile(path, "r") as zf:
            xml_bytes = zf.read("word/document.xml")
        root = ET.fromstring(xml_bytes)
        texts = [node.text for node in root.iter() if node.tag.endswith("}t") and node.text]
        text = " ".join(texts)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=None, status=status)
    except Exception as exc:
        logger.warning("Could not read DOCX %s: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")


def _ingest_xlsx(path: Path) -> NormalizedDocument:
    """Extract sheet names and cell text from an Excel (.xlsx) workbook."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            lines: list[str] = []
            sheet_count = len(wb.sheetnames)
            for sheet_name in wb.sheetnames:
                lines.append(f"[Sheet: {sheet_name}]")
                ws = wb[sheet_name]
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals:
                        lines.append(", ".join(vals))
            text = "\n".join(lines)
            has_cells = any(not line.startswith("[Sheet:") for line in lines)
            status = "ok" if has_cells else "empty"
            return _make_doc(path, text=text if has_cells else "", page_count=sheet_count, status=status)
        finally:
            wb.close()
    except ImportError:
        pass
    except Exception as exc:
        logger.debug("openpyxl failed on %s (%s), trying XML fallback", path, exc)

    # Stdlib zipfile + XML fallback for .xlsx
    try:
        with zipfile.ZipFile(path, "r") as zf:
            strings: list[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                strings.extend(n.text for n in root.iter() if n.tag.endswith("}t") and n.text)
            sheet_files = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
            for sf in sheet_files:
                root = ET.fromstring(zf.read(sf))
                strings.extend(n.text for n in root.iter() if n.tag.endswith("}v") and n.text)
        text = " ".join(strings)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=len(sheet_files) or None, status=status)
    except Exception as exc:
        logger.warning("Could not read XLSX %s: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")


def _ingest_pptx(path: Path) -> NormalizedDocument:
    """Extract slide text from a PowerPoint (.pptx) presentation."""
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides_text: list[str] = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    txt = shape.text_frame.text.strip()
                    if txt:
                        slide_parts.append(txt)
            if slide_parts:
                slides_text.append(f"[Slide {idx}] " + " ".join(slide_parts))
        text = "\n".join(slides_text)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=len(prs.slides), status=status)
    except ImportError:
        pass
    except Exception as exc:
        logger.debug("python-pptx failed on %s (%s), trying XML fallback", path, exc)

    # Stdlib zipfile + XML fallback for .pptx
    try:
        with zipfile.ZipFile(path, "r") as zf:
            slide_files = sorted(
                n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")
            )
            slides_text = []
            for sf in slide_files:
                root = ET.fromstring(zf.read(sf))
                parts = [n.text for n in root.iter() if n.tag.endswith("}t") and n.text]
                if parts:
                    slides_text.append(" ".join(parts))
        text = "\n".join(slides_text)
        status = "ok" if text.strip() else "empty"
        return _make_doc(path, text=text, page_count=len(slide_files), status=status)
    except Exception as exc:
        logger.warning("Could not read PPTX %s: %s", path, exc)
        return _make_doc(path, text="", page_count=None, status="error")


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def ingest_file(path: str | Path) -> NormalizedDocument:
    """
    Ingest *path* and return a NormalizedDocument.

    Never raises for recoverable errors (missing file, bad PDF, unsupported
    format); those are captured in extraction_status instead.
    """
    path = Path(path).resolve()

    if not path.exists():
        logger.warning("File not found: %s", path)
        return NormalizedDocument(
            **_base_fields(path),
            text="",
            character_count=0,
            page_count=None,
            extraction_status="missing",
        )

    ext = path.suffix.lstrip(".").lower()

    if ext == "pdf":
        return _ingest_pdf(path)

    if ext in _TEXT_EXTENSIONS:
        return _ingest_text(path)

    if ext in _IMAGE_EXTENSIONS:
        return _ingest_image(path)

    if ext == "docx":
        return _ingest_docx(path)

    if ext == "xlsx":
        return _ingest_xlsx(path)

    if ext == "pptx":
        return _ingest_pptx(path)

    logger.info("Unsupported extension '%s' for file: %s", ext, path)
    return _make_doc(path, text="", page_count=None, status="unsupported")

